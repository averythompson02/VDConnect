# -*- coding: utf-8 -*-

from pyrevit import revit, DB, forms, script
import openpyxl
import math

doc = revit.doc
output = script.get_output()

# -------------------------------
# SELECT EXCEL FILE
# -------------------------------

excel_file = forms.pick_file(
    file_ext='xlsx',
    title='Select Excel File'
)

if not excel_file:
    script.exit()

wb = openpyxl.load_workbook(
    excel_file,
    data_only=True
)

sheet_name = forms.SelectFromList.show(
    wb.sheetnames,
    title='Select Worksheet',
    multiselect=False
)

if not sheet_name:
    script.exit()

ws = wb[sheet_name]

# -------------------------------
# READ HEADERS
# -------------------------------

headers = []

for cell in ws[1\]:
    headers.append(str(cell.value))

# -------------------------------
# COLUMN MAPPING
# -------------------------------

x_col = forms.SelectFromList.show(
    headers,
    title='X Coordinate Column'
)

y_col = forms.SelectFromList.show(
    headers,
    title='Y Coordinate Column'
)

z_col = forms.SelectFromList.show(
    ['<None>'] + headers,
    title='Z Coordinate Column'
)

rot_col = forms.SelectFromList.show(
    ['<None>'] + headers,
    title='Rotation Column'
)

id_col = forms.SelectFromList.show(
    ['<None>'] + headers,
    title='Point ID Column'
)

# -------------------------------
# CATEGORY PICKER
# -------------------------------

category_options = [
    DB.BuiltInCategory.OST_GenericModel,
    DB.BuiltInCategory.OST_MechanicalEquipment,
    DB.BuiltInCategory.OST_PlumbingFixtures,
    DB.BuiltInCategory.OST_SpecialityEquipment
]

category_names = {}

for bic in category_options:
    category_names[str(bic)] = bic

selected_category_name = forms.SelectFromList.show(
    sorted(category_names.keys()),
    title='Select Category'
)

if not selected_category_name:
    script.exit()

selected_category = category_names[selected_category_name]

# -------------------------------
# FAMILY PICKER
# -------------------------------

symbols = list(
    DB.FilteredElementCollector(doc)
    .OfClass(DB.FamilySymbol)
)

family_dict = {}

for symbol in symbols:

    try:

        category = symbol.Category

        if not category:
            continue

        if category.Id.IntegerValue != int(selected_category):
            continue

        family_name = symbol.Family.Name

        if family_name not in family_dict:
            family_dict[family_name] = []

        family_dict[family_name].append(symbol)

    except:
        pass

if not family_dict:
    forms.alert(
        'No family symbols found.',
        exitscript=True
    )

selected_family = forms.SelectFromList.show(
    sorted(family_dict.keys()),
    title='Select Family'
)

if not selected_family:
    script.exit()

# -------------------------------
# TYPE PICKER
# -------------------------------

family_symbols = sorted(
    family_dict[selected_family],
    key=lambda x: x.Name
)

type_names = [x.Name for x in family_symbols]

selected_type_name = forms.SelectFromList.show(
    type_names,
    title='Select Type'
)

if not selected_type_name:
    script.exit()

selected_symbol = None

for symbol in family_symbols:
    if symbol.Name == selected_type_name:
        selected_symbol = symbol
        break

# -------------------------------
# COORDINATE SYSTEM
# -------------------------------

coordinate_system = forms.SelectFromList.show(
    [
        "Shared Coordinates",
        "Internal Coordinates"
    ],
    title='Coordinate System'
)

# -------------------------------
# HEADER INDEXES
# -------------------------------

x_index = headers.index(x_col)
y_index = headers.index(y_col)

z_index = headers.index(z_col) if z_col != '<None>' else None
rot_index = headers.index(rot_col) if rot_col != '<None>' else None
id_index = headers.index(id_col) if id_col != '<None>' else None

# -------------------------------
# ACTIVATE FAMILY
# -------------------------------

t = DB.Transaction(
    doc,
    "Place Family From Excel"
)

t.Start()

if not selected_symbol.IsActive:
    selected_symbol.Activate()
    doc.Regenerate()

placed = 0
errors = []

# -------------------------------
# CREATE INSTANCES
# -------------------------------

for row in ws.iter_rows(min_row=2):

    try:

        x = float(row[x_index].value)
        y = float(row[y_index].value)

        if z_index is not None:
            z = float(row[z_index].value)
        else:
            z = 0.0

        point = DB.XYZ(
            x,
            y,
            z
        )

        # Shared to Internal
        if coordinate_system == 'Shared Coordinates':

            transform = (
                doc.ActiveProjectLocation
                .GetTotalTransform()
            )

            point = transform.Inverse.OfPoint(
                point
            )

        element = doc.Create.NewFamilyInstance(
            point,
            selected_symbol,
            DB.Structure.StructuralType.NonStructural
        )

        # Rotation
        if rot_index is not None:

            rotation = float(
                row[rot_index].value
            )

            axis = DB.Line.CreateBound(
                point,
                point + DB.XYZ.BasisZ
            )

            DB.ElementTransformUtils.RotateElement(
                doc,
                element.Id,
                axis,
                math.radians(rotation)
            )

        # Point ID
        if id_index is not None:

            point_id = str(
                row[id_index].value
            )

            comments = element.get_Parameter(
                DB.BuiltInParameter.ALL_MODEL_INSTANCE_COMMENTS
            )

            if comments and not comments.IsReadOnly:
                comments.Set(point_id)

        placed += 1

    except Exception as ex:

        errors.append(
            str(ex)
        )

t.Commit()

# -------------------------------
# REPORT
# -------------------------------

output.print_md(
    "## Placement Complete"
)

output.print_md(
    "**Family:** {}".format(
        selected_family
    )
)

output.print_md(
    "**Type:** {}".format(
        selected_type_name
    )
)

output.print_md(
    "**Placed:** {
        placed
    )
)

output.print_md(
    "**Errors:** {}".format(
        len(errors)
    )
)

if errors:

    output.print_md(
        "## Errors"
    )

    for err in errors:
        output.print_md(
            "- {}".format(err)
        )

forms.alert(
    "{} families placed".format(
        placed
    )
)